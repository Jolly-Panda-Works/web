#!/usr/bin/env python3
"""
Jolly Panda Web — make-brochure.py
Builds the downloadable price-list brochures (PDF, Persian + English) from the PowerPoint
template in brochure/ using the CURRENT prices:

    data/plans.json  (regular prices in USD, content, discounts)  +  data/rate.json (USD -> IRR)

  downloads/Jolly-Panda-Web-Price-List-fa.pdf   Persian, right-to-left, prices in Rial (USD if no rate yet)
  downloads/Jolly-Panda-Web-Price-List-en.pdf   English, mirrored left-to-right, prices in US dollars

Prices follow exactly the same rules as the website (js/pricing.js): the discount is taken from the row,
then the type, then the plan, then the global default; discounted = round(regular x (100 - pct)) / 100.

Usage:   python3 scripts/make-brochure.py [--keep-pptx]
Needs:   pip install python-pptx   and LibreOffice (soffice) for the PDF export.
         SOFFICE="<command>" overrides the LibreOffice command.
The GitHub Action (.github/workflows/update-rate.yml) runs it after the daily rate update
and whenever data/plans.json or the template changes; the PDFs are committed to downloads/.
"""
import copy, json, math, os, re, shlex, shutil, subprocess, sys, tempfile
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from lxml import etree
from pptx import Presentation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
TEMPLATE = os.path.join(ROOT, "brochure", "JollyPanda_Website_Plans_template.pptx")
OUT_DIR = os.path.join(ROOT, "downloads")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}
EMU = 914400

# ------------------------------------------------------------------ numbers, dates
FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
def half_up(x): return int(math.floor(x + 0.5))          # same rounding as JavaScript Math.round for positive numbers

def num(n, lang):
    """Whole numbers with thousands separators; amounts with cents (24.99) always show two decimals."""
    s = f"{int(n):,}" if float(n) == int(n) else f"{n:,.2f}"
    if lang == "fa": s = s.replace(",", "٬").replace(".", "٫").translate(FA_DIGITS)
    return s

def g2j(gy, gm, gd):                                      # Gregorian -> Jalali
    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    if gy > 1600: jy, gy = 979, gy - 1600
    else: jy, gy = 0, gy - 621
    gy2 = gy + 1 if gm > 2 else gy
    days = 365 * gy + (gy2 + 3) // 4 - (gy2 + 99) // 100 + (gy2 + 399) // 400 - 80 + gd + g_d_m[gm - 1]
    jy += 33 * (days // 12053); days %= 12053
    jy += 4 * (days // 1461); days %= 1461
    if days > 365: jy += (days - 1) // 365; days = (days - 1) % 365
    if days < 186: jm, jd = 1 + days // 31, 1 + days % 31
    else: jm, jd = 7 + (days - 186) // 30, 1 + (days - 186) % 30
    return jy, jm, jd

FA_MONTHS = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
EN_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

def date_text(iso, lang):
    d = datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(ZoneInfo("Asia/Tehran"))
    if lang == "fa":
        jy, jm, jd = g2j(d.year, d.month, d.day)
        return f"{jd} {FA_MONTHS[jm - 1]} {jy}".translate(FA_DIGITS)
    return f"{d.day} {EN_MONTHS[d.month - 1]} {d.year}"

# ------------------------------------------------------------------ data + price rules (mirror of js/pricing.js)
plans = json.load(open(os.path.join(ROOT, "data", "plans.json"), encoding="utf-8"))
rate_file = json.load(open(os.path.join(ROOT, "data", "rate.json"), encoding="utf-8"))
RATE = rate_file.get("rateRial") if isinstance(rate_file.get("rateRial"), (int, float)) and rate_file.get("rateRial", 0) > 0 else None
UPDATED_ISO = (rate_file.get("fetchedAt") if RATE else None) or plans.get("listUpdated")

def clean(p):
    try: p = float(p)
    except (TypeError, ValueError): return 0
    return p if 0 < p < 100 else 0

def pct_for(ty, pl):
    for v in (ty["rows"][pl["id"]].get("discount"), ty.get("discount"), pl.get("discount"), plans.get("discountPercent")):
        if v is not None: return clean(v)
    return 0

def discounted(usd, p): return half_up(usd * (100 - clean(p))) / 100          # to the cent (same as js/pricing.js)

ALL_PCTS = [pct_for(t, p) for t in plans["types"] for p in plans["plans"]]
MAX_PCT, MIN_PCT = max(ALL_PCTS), min(ALL_PCTS)

def money(usd, lang):
    """Price text: Persian -> Rial when a rate exists (else dollars), English -> dollars."""
    if lang == "fa":
        if RATE: return f"{num(half_up(usd * RATE), 'fa')} ریال"
        return f"{num(usd, 'fa')} دلار"
    return f"${num(usd, 'en')}"

# ------------------------------------------------------------------ static texts (Persian = the template's own wording)
T = {
  "fa": {
    "brand": "Jolly Panda Studio",
    "cover_title": "پلن‌ها و تعرفه‌های طراحی و پیاده‌سازی وبسایت",
    "cover_types": "بلاگ · پورتفولیو · شرکتی · آموزشگاهی · خبری · فروشگاهی",
    "cover_plans": "سه پلن اکونومی، ویژه و اختصاصی برای هر نوع وبسایت",
    "cmp_title": "مقایسه قیمت پلن‌ها",
    "cmp_sub_rial": "قیمت شروع هر پلن بر حسب ریال (با نرخ روز دلار)؛ جزئیات هر نوع وبسایت در اسلایدهای بعد",
    "cmp_sub_usd": "قیمت شروع هر پلن بر حسب دلار آمریکا؛ جزئیات هر نوع وبسایت در اسلایدهای بعد",
    "type_hdr": "نوع وبسایت", "feat_hdr": "ویژگی", "price_row": "قیمت شروع",
    "badge_special": "پیشنهاد ویژه",
    "off": "{p}٪ تخفیف", "off_all": "{p}٪ تخفیف روی همه پلن‌ها", "off_up": "تا {p}٪ تخفیف",
    "updated": "آخرین بروزرسانی قیمت: {d}",
    "end_title": "آماده شروع هستید؟",
    "end_sub": "قیمت نهایی در جلسه مشاوره با تیم Jolly Panda تعیین می‌شود",
    "end_1": "قیمت‌ها نقطه شروع هر پلن هستند", "end_1_off": "قیمت‌ها نقطه شروع هر پلن هستند و تخفیف روی آن‌ها اعمال شده است",
    "end_2": "هزینه دامنه و هاست جدا از پلن‌ها محاسبه می‌شود",
    "end_3": "همه پلن‌ها شامل طراحی واکنش‌گرا (موبایل و تبلت) و گواهی SSL هستند",
  },
  "en": {
    "brand": "Jolly Panda Studio",
    "cover_title": "Website design & development plans and prices",
    "cover_types": "Blog · Portfolio · Corporate · Educational · News · Online store",
    "cover_plans": "Three plans — Economy, Special and Custom — for every website type",
    "cmp_title": "Plan price comparison",
    "cmp_sub_rial": "", "cmp_sub_usd": "Starting price of each plan in US dollars; details for each website type on the following slides",
    "type_hdr": "Website type", "feat_hdr": "Feature", "price_row": "Starting price",
    "badge_special": "Special offer",
    "off": "{p}% off", "off_all": "{p}% off all plans", "off_up": "Up to {p}% off",
    "updated": "Prices last updated: {d}",
    "end_title": "Ready to get started?",
    "end_sub": "The final price is set in a consultation with the Jolly Panda team",
    "end_1": "Prices are the starting point of each plan", "end_1_off": "Prices are the starting point of each plan, with the discount already applied",
    "end_2": "Domain and hosting costs are charged separately from the plans",
    "end_3": "All plans include responsive design (mobile and tablet) and an SSL certificate",
  },
}

# ------------------------------------------------------------------ low-level helpers (lxml on python-pptx elements)
def q(tag): p_, t = tag.split(":"); return "{%s}%s" % (NS[p_], t)

def set_para_runs(p_el, pieces):
    """Replace the runs of a paragraph. pieces = [(text, {sz, strike, color, bold})]; formatting is cloned from the first run."""
    runs = p_el.findall(q("a:r"))
    base = copy.deepcopy(runs[0])
    for r in runs: p_el.remove(r)
    end = p_el.find(q("a:endParaRPr"))
    for text, o in pieces:
        r = copy.deepcopy(base)
        rpr = r.find(q("a:rPr"))
        if "sz" in o: rpr.set("sz", str(int(round(o["sz"] * 100))))
        if "bold" in o: rpr.set("b", "1" if o["bold"] else "0")
        if o.get("strike"): rpr.set("strike", "sngStrike")
        if "color" in o:
            clr = rpr.find(q("a:solidFill") + "/" + q("a:srgbClr"))
            if clr is not None: clr.set("val", o["color"])
        r.find(q("a:t")).text = text
        if end is not None: end.addprevious(r)
        else: p_el.append(r)

def set_text(txbody_el, text):
    """Plain text into the first paragraph of a txBody (formatting kept); extra paragraphs are removed."""
    paras = txbody_el.findall(q("a:p"))
    for extra in paras[1:]: txbody_el.remove(extra)
    set_para_runs(paras[0], [(text, {})])

def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name: return sh
    raise KeyError(name)

def run_size(p_el):
    r = p_el.find(q("a:r") + "/" + q("a:rPr"))
    return int(r.get("sz")) / 100 if r is not None and r.get("sz") else 14

def set_price_cell(tc, usd, pct, lang, base_size):
    """Price cell: discounted price (bold) + a small struck-through regular price and the discount label above it."""
    txb = tc.find(q("a:txBody"))
    paras = txb.findall(q("a:p"))
    for extra in paras[1:]: txb.remove(extra)
    main_p = paras[0]
    now = money(discounted(usd, pct), lang)
    long_text = lang == "fa" and RATE                       # Rial amounts are long: smaller type
    size = min(base_size, 18 if base_size >= 20 else base_size) if long_text else base_size
    if not pct:
        set_para_runs(main_p, [(now, {"sz": size})])
        return
    small = max(9.5, round(size * 0.55, 1))
    top = copy.deepcopy(main_p)
    set_para_runs(top, [(money(usd, lang), {"sz": small, "strike": True, "color": "8A7268", "bold": False}),
                        ("   ", {"sz": small}),
                        (T[lang]["off"].format(p=num(pct, lang)), {"sz": small, "color": "C4570D", "bold": True})])
    set_para_runs(main_p, [(now, {"sz": size})])
    main_p.addprevious(top)

# ------------------------------------------------------------------ mirroring (English)
def mirror_slide(slide, slide_w):
    for sh in slide.shapes:
        sh.left = slide_w - sh.left - sh.width
    for p_el in slide._element.iter(q("a:p")):
        ppr = p_el.find(q("a:pPr"))
        if ppr is not None and ppr.get("rtl") == "1": ppr.set("rtl", "0")
    for el in slide._element.iter():
        if el.tag in (q("a:rPr"), q("a:endParaRPr")):
            if el.get("lang", "").startswith("fa"): el.set("lang", "en-US")
            if "altLang" in el.attrib: del el.attrib["altLang"]

def flip_alignment(slide, mirrored_left_half):
    """Right-aligned text becomes left-aligned; default (left) text of shapes that moved to the right becomes right-aligned."""
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip(): continue
        for p_el in sh._element.iter(q("a:p")):
            ppr = p_el.find(q("a:pPr"))
            if ppr is None: ppr = etree.SubElement(p_el, q("a:pPr")); p_el.insert(0, ppr)
            a = ppr.get("algn")
            if a == "r": ppr.set("algn", "l")
            elif a in (None, "l") and sh.shape_id in mirrored_left_half: ppr.set("algn", "r")

def mirror_table(gf):
    tbl = gf._element.find(".//" + q("a:tbl"))
    grid = tbl.find(q("a:tblGrid"))
    cols = list(grid); [grid.remove(c) for c in cols]; [grid.append(c) for c in reversed(cols)]
    for tr in tbl.findall(q("a:tr")):
        tcs = tr.findall(q("a:tc")); [tr.remove(c) for c in tcs]
        for c in reversed(tcs): tr.append(c)
        for p_el in tr.iter(q("a:p")):
            ppr = p_el.find(q("a:pPr"))
            if ppr is not None:
                if ppr.get("rtl") == "1": ppr.set("rtl", "0")
                if ppr.get("algn") == "r": ppr.set("algn", "l")
        for el in tr.iter():
            if el.tag in (q("a:rPr"), q("a:endParaRPr")):
                if el.get("lang", "").startswith("fa"): el.set("lang", "en-US")
                if "altLang" in el.attrib: del el.attrib["altLang"]

def next_id(slide): return max(int(e.get("id")) for e in slide._element.iter(q("p:cNvPr"))) + 1

def clone_shape(slide, name, new_name, **kw):
    sh = shape_by_name(slide, name)
    el = copy.deepcopy(sh._element)
    sh._element.addnext(el)
    c = el.find(".//" + q("p:cNvPr")); c.set("id", str(next_id(slide))); c.set("name", new_name)
    new = [s for s in slide.shapes if s.name == new_name][-1]
    for k, v in kw.items(): setattr(new, k, v)
    return new

# ------------------------------------------------------------------ build one language
def build(lang):
    prs = Presentation(TEMPLATE)
    W = prs.slide_width
    t = T[lang]
    slides = list(prs.slides)
    PLAN_ORDER = ["custom", "special", "economy"]               # left -> right in the template (right-to-left design)
    plan_by_id = {p["id"]: p for p in plans["plans"]}
    pick = lambda f: f[lang]
    date_str = t["updated"].format(d=date_text(UPDATED_ISO, lang)) if UPDATED_ISO else ""

    # badge text: same rule as the site banner
    if MAX_PCT > 0:
        discount_badge = t["off_all"].format(p=num(MAX_PCT, lang)) if MIN_PCT == MAX_PCT else t["off_up"].format(p=num(MAX_PCT, lang))
    else:
        discount_badge = ""

    def footer_date(slide):
        if not date_str: return
        src = shape_by_name(slide, "Text 4")           # "www.jollypanda.ir" footer text
        new = clone_shape(slide, "Text 4", "Date footer", left=int((W - 5.2 * EMU) / 2), width=int(5.2 * EMU))
        for p_el in new._element.iter(q("a:p")):
            ppr = p_el.find(q("a:pPr"))
            if ppr is None: ppr = etree.Element(q("a:pPr")); p_el.insert(0, ppr)
            ppr.set("algn", "ctr")
        set_text(new._element.find(q("p:txBody")), date_str)
        return new

    def discount_badge_shapes(slide, shape_name, text_name, right_edge_in, width_in=2.7):
        """Yellow label cloned from the 'Special offer' badge, right-aligned with the table's right edge."""
        if not discount_badge: return
        center_x_in = right_edge_in - width_in / 2
        bg = clone_shape(slide, shape_name, "Discount badge", left=int((center_x_in - width_in / 2) * EMU), width=int(width_in * EMU))
        fill = bg._element.find(".//" + q("a:solidFill") + "/" + q("a:srgbClr"))
        if fill is not None: fill.set("val", "FCDD4E")
        tx = clone_shape(slide, text_name, "Discount badge text", left=int((center_x_in - width_in / 2) * EMU), width=int(width_in * EMU))
        for clr in tx._element.iter(q("a:srgbClr")): clr.set("val", "3A2415")
        set_text(tx._element.find(q("p:txBody")), discount_badge)

    # ---- slide 1: cover
    s = slides[0]
    set_text(shape_by_name(s, "Text 2")._element.find(q("p:txBody")), t["cover_title"])
    types_p = shape_by_name(s, "Text 3")._element.find(q("p:txBody")).find(q("a:p"))
    set_para_runs(types_p, [(t["cover_types"], {"sz": run_size(types_p) - (1.5 if lang == "en" else 0)})])   # the English list is longer
    set_text(shape_by_name(s, "Text 4")._element.find(q("p:txBody")), t["cover_plans"])

    # ---- slide 2: comparison table
    s = slides[1]
    set_text(shape_by_name(s, "Text 0")._element.find(q("p:txBody")), t["cmp_title"])
    set_text(shape_by_name(s, "Text 1")._element.find(q("p:txBody")), t["cmp_sub_rial"] if (lang == "fa" and RATE) else t["cmp_sub_usd"])
    set_text(shape_by_name(s, "Text 3")._element.find(q("p:txBody")), t["badge_special"])
    gf = shape_by_name(s, "Table 0"); tbl = gf._element.find(".//" + q("a:tbl"))
    rows = tbl.findall(q("a:tr"))
    hdr = rows[0].findall(q("a:tc"))
    for i, pid in enumerate(PLAN_ORDER): set_text(hdr[i].find(q("a:txBody")), pick(plan_by_id[pid]["name"]))
    set_text(hdr[3].find(q("a:txBody")), t["type_hdr"])
    for r_i, ty in enumerate(plans["types"]):
        tcs = rows[r_i + 1].findall(q("a:tc"))
        for i, pid in enumerate(PLAN_ORDER):
            base = run_size(tcs[i].find(q("a:txBody")).find(q("a:p")))
            set_price_cell(tcs[i], ty["rows"][pid]["price"], pct_for(ty, plan_by_id[pid]), lang, base)
        set_text(tcs[3].find(q("a:txBody")), pick(ty["name"]))
    for tr in rows[1:]: tr.set("h", str(int(0.62 * EMU)))        # room for the struck-through line
    discount_badge_shapes(s, "Shape 2", "Text 3", (gf.left + gf.width) / EMU)
    footer_date(s)

    # ---- slides 3..8: one per website type
    for idx, ty in enumerate(plans["types"]):
        s = slides[2 + idx]
        set_text(shape_by_name(s, "Text 0")._element.find(q("p:txBody")), pick(ty["title"]))
        set_text(shape_by_name(s, "Text 1")._element.find(q("p:txBody")), pick(ty["tagline"]))
        set_text(shape_by_name(s, "Text 3")._element.find(q("p:txBody")), t["badge_special"])
        gf = shape_by_name(s, "Table 0"); tbl = gf._element.find(".//" + q("a:tbl"))
        rows = tbl.findall(q("a:tr"))
        hdr = rows[0].findall(q("a:tc"))
        for i, pid in enumerate(PLAN_ORDER): set_text(hdr[i].find(q("a:txBody")), pick(plan_by_id[pid]["name"]))
        set_text(hdr[3].find(q("a:txBody")), t["feat_hdr"])
        for r_i, ft in enumerate(plans["features"]):
            tcs = rows[r_i + 1].findall(q("a:tc"))
            label = ty["structureLabel"] if ft["id"] == "structure" and ty.get("structureLabel") else ft["label"]
            for i, pid in enumerate(PLAN_ORDER): set_text(tcs[i].find(q("a:txBody")), pick(ty["rows"][pid][ft["id"]]))
            set_text(tcs[3].find(q("a:txBody")), pick(label))
        last = rows[-1].findall(q("a:tc"))
        for i, pid in enumerate(PLAN_ORDER):
            base = run_size(last[i].find(q("a:txBody")).find(q("a:p")))
            set_price_cell(last[i], ty["rows"][pid]["price"], pct_for(ty, plan_by_id[pid]), lang, base)
        set_text(last[3].find(q("a:txBody")), t["price_row"])
        rows[-1].set("h", str(int(0.88 * EMU)))
        discount_badge_shapes(s, "Shape 2", "Text 3", (gf.left + gf.width) / EMU)
        footer_date(s)

    # ---- slide 9: closing
    s = slides[8]
    set_text(shape_by_name(s, "Text 1")._element.find(q("p:txBody")), t["end_title"])
    set_text(shape_by_name(s, "Text 2")._element.find(q("p:txBody")), t["end_sub"])
    set_text(shape_by_name(s, "Text 4")._element.find(q("p:txBody")), t["end_1_off"] if MAX_PCT > 0 else t["end_1"])
    set_text(shape_by_name(s, "Text 6")._element.find(q("p:txBody")), t["end_2"])
    set_text(shape_by_name(s, "Text 8")._element.find(q("p:txBody")), t["end_3"])
    if date_str:
        d = clone_shape(s, "Text 11", "Date footer", left=int((W - 5.2 * EMU) / 2), width=int(5.2 * EMU))
        for p_el in d._element.iter(q("a:p")):
            ppr = p_el.find(q("a:pPr"))
            if ppr is None: ppr = etree.Element(q("a:pPr")); p_el.insert(0, ppr)
            ppr.set("algn", "ctr")
        set_text(d._element.find(q("p:txBody")), date_str)

    # ---- English: mirror everything to left-to-right
    if lang == "en":
        for s in slides:
            left_half = {sh.shape_id for sh in s.shapes if (sh.left + sh.width / 2) < W / 2}
            for sh in s.shapes:
                if sh.has_table: mirror_table(sh)
            flip_alignment(s, left_half)       # decided on the ORIGINAL positions
            mirror_slide(s, W)

    prs.core_properties.title = "Jolly Panda Web — " + t["cover_title"]
    prs.core_properties.author = "Jolly Panda Studio"
    out = os.path.join(WORK_DIR, f"Jolly-Panda-Web-Price-List-{lang}.pptx")
    prs.save(out)
    return out

def to_pdf(pptx):
    cmd = shlex.split(os.environ.get("SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice") or "")
    if not cmd: sys.exit("LibreOffice not found (install libreoffice-impress or set SOFFICE)")
    subprocess.run(cmd + ["--headless", "--convert-to", "pdf", "--outdir", WORK_DIR, pptx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.splitext(pptx)[0] + ".pdf"
    if not os.path.exists(pdf): sys.exit("PDF export failed for " + pptx)
    return pdf

if __name__ == "__main__":
    keep = "--keep-pptx" in sys.argv
    WORK_DIR = tempfile.mkdtemp(prefix="jp-brochure-")
    os.makedirs(OUT_DIR, exist_ok=True)
    for lang in ("fa", "en"):
        pptx = build(lang)
        pdf = to_pdf(pptx)
        shutil.copyfile(pdf, os.path.join(OUT_DIR, os.path.basename(pdf)))
        if keep: shutil.copyfile(pptx, os.path.join(OUT_DIR, os.path.basename(pptx)))
        print("wrote", os.path.join("downloads", os.path.basename(pdf)))
    shutil.rmtree(WORK_DIR, ignore_errors=True)
    print(f"rate: {RATE}  updated: {UPDATED_ISO}  discounts: {MIN_PCT}-{MAX_PCT}%")
